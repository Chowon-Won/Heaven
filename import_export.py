"""Lectura y escritura de CSV y Excel sin perder ceros iniciales."""
from __future__ import annotations

import csv
import io
import re
import unicodedata
from pathlib import Path

HEADER_ALIASES = {
    "codigo": "codigo", "cod": "codigo", "code": "codigo",
    "codigo escuela": "codigo", "codigo centro": "codigo",
    "codigo producto": "codigo",
    "nombre": "nombre", "name": "nombre", "nombre completo": "nombre",
    "nombre centro": "nombre", "nombre centro o cliente": "nombre",
    "escuela": "nombre", "producto": "nombre",
    "orden": "orden", "orden ruta": "orden", "orden en la ruta": "orden", "no": "orden",
    "cantidad": "cantidad", "cant": "cantidad", "raciones": "cantidad",
    "cant raciones": "cantidad", "cantidad raciones": "cantidad",
    "activo": "activo", "estado": "activo", "activa": "activo",
    "notas": "notas", "nota": "notas", "observaciones": "notas",
    "descripcion": "descripcion",
    "texto verificacion": "texto_verificacion",
    "texto_verificacion": "texto_verificacion",
    "descripcion sisfacon": "texto_verificacion",
}


def _norm_header(h) -> str:
    text = unicodedata.normalize("NFKD", str(h or "")).encode("ascii", "ignore").decode()
    text = re.sub(r"[^a-z0-9_ ]", " ", text.lower()).strip()
    text = re.sub(r"\s+", " ", text)
    return HEADER_ALIASES.get(text, text.replace(" ", "_"))


def read_table(path: str | Path) -> list[dict]:
    path = Path(path)
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        return _read_xlsx(path)
    return _read_csv(path)


def _read_csv(path: Path) -> list[dict]:
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ";" if sample.count(";") > sample.count(",") else ","
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    rows = list(reader)
    if not rows:
        return []
    headers = [_norm_header(h) for h in rows[0]]
    # Todo se lee como texto: los ceros iniciales se conservan.
    return [dict(zip(headers, (c.strip() for c in r))) for r in rows[1:]]


def _read_xlsx(path: Path) -> list[dict]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=False, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows())
    if not rows:
        return []
    headers = [_norm_header(c.value) for c in rows[0]]
    out = []
    for r in rows[1:]:
        item = {}
        for h, cell in zip(headers, r):
            item[h] = _cell_text(cell)
        out.append(item)
    return out


def _cell_text(cell) -> str:
    """Devuelve el texto de una celda respetando formatos como «00000»."""
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if isinstance(value, float) and value.is_integer():
            value = int(value)
        text = str(value)
        fmt = str(cell.number_format or "")
        if re.fullmatch(r"0+", fmt) and len(fmt) > len(text):
            text = text.zfill(len(fmt))
        return text
    return str(value).strip()


def write_table(path: str | Path, headers: list[str], rows: list[list]) -> None:
    path = Path(path)
    if path.suffix.lower() == ".xlsx":
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.append(headers)
        for r in rows:
            ws.append(["" if v is None else str(v) for v in r])
        # Columnas como texto para que Excel no elimine los ceros.
        for col in ws.iter_cols(min_row=2):
            for cell in col:
                cell.number_format = "@"
        wb.save(path)
        return
    with open(path, "w", newline="", encoding="utf-8-sig") as fh:
        writer = csv.writer(fh, delimiter=";")
        writer.writerow(headers)
        for r in rows:
            writer.writerow(["" if v is None else v for v in r])
